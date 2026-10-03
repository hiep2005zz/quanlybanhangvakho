import io
<<<<<<< HEAD
import re
from typing import List, Dict, Any, Tuple
import openpyxl
from pydantic import ValidationError
from app.models.user import USERS_DB, UserInDB, get_next_user_id, save_users_db
from app.schemas.user_import import RowPreview, PreviewResponse, ImportExecuteRequest, ImportExecuteResponse
from app.core.rbac import Role, is_warehouse_role, is_specific_warehouse, SPECIFIC_WAREHOUSES
from app.core.security import get_password_hash
from app.schemas.user import UserCreate

def generate_temp_password(length: int = 8) -> str:
    import secrets
    import string
    alphabet = string.ascii_letters + string.digits
    return ''.join(secrets.choice(alphabet) for i in range(length))

def validate_excel_file(file_content: bytes) -> PreviewResponse:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
        sheet = wb.active
    except Exception as e:
        raise ValueError("Không thể đọc tệp Excel. Vui lòng đảm bảo tệp đúng định dạng .xlsx")

    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) < 2:
        return PreviewResponse(total_rows=0, valid_rows=0, invalid_rows=0, preview_data=[])

    header = rows[0]
    data_rows = rows[1:]

    preview_data: List[RowPreview] = []
    
    seen_emails = set()
    seen_phones = set()

    for idx, row in enumerate(data_rows):
        row_index = idx + 2
        if not any(row):
            continue

        full_name = str(row[0] or "").strip()
        email = str(row[1] or "").strip()
        phone = str(row[2] or "").strip()
        if phone.endswith(".0"):
            phone = phone[:-2]
            
        role = str(row[3] or "").strip()
        branch = str(row[4] or "").strip()
        password = str(row[5] or "").strip()

        errors = []

        if not full_name:
            errors.append("Thiếu Họ và tên")

        if not email:
            errors.append("Thiếu Email")
        else:
            email_lower = email.lower()
            if not re.match(r"[^@]+@[^@]+\.[^@]+", email_lower):
                errors.append("Email không đúng định dạng")
            elif email_lower in seen_emails:
                errors.append("Email trùng lặp trong tệp Excel")
            else:
                seen_emails.add(email_lower)
                for u in USERS_DB.values():
                    if u.email and u.email.lower() == email_lower:
                        errors.append("Email đã tồn tại trong hệ thống")
                        break

        if phone:
            cleaned_phone = re.sub(r'[\s\.\-\(\)]', '', phone)
            if cleaned_phone.startswith('+84'):
                cleaned_phone = '0' + cleaned_phone[3:]
            elif cleaned_phone.startswith('84') and len(cleaned_phone) == 11:
                cleaned_phone = '0' + cleaned_phone[2:]
            
            vn_phone_pattern = r'^(0[3|5|7|8|9][0-9]{8}|02[0-9]{9})$'
            if not re.match(vn_phone_pattern, cleaned_phone):
                errors.append("Số điện thoại không đúng định dạng (cần 10 số)")
            elif cleaned_phone in seen_phones:
                errors.append("Số điện thoại trùng lặp trong tệp Excel")
            else:
                seen_phones.add(cleaned_phone)
                phone = cleaned_phone

        valid_roles = [r.value for r in Role]
        if not role:
            errors.append("Thiếu Vai trò")
        elif role not in valid_roles:
            errors.append(f"Vai trò không hợp lệ (hợp lệ: {', '.join(valid_roles)})")

        if role and role in valid_roles:
            if is_warehouse_role([role]):
                if not branch:
                    errors.append("Vai trò Kho yêu cầu chỉ định Phòng ban/Địa bàn")
                elif not is_specific_warehouse(branch):
                    errors.append(f"Địa bàn Kho không hợp lệ (Ví dụ: {', '.join(SPECIFIC_WAREHOUSES)})")

        if not password:
            password = generate_temp_password()

        is_valid = len(errors) == 0

        preview = RowPreview(
            row_index=row_index,
            full_name=full_name,
            email=email,
            phone=phone,
            role=role,
            branch=branch,
            password=password,
            is_valid=is_valid,
            errors=errors
        )
        preview_data.append(preview)

    total_rows = len(preview_data)
    valid_rows = sum(1 for r in preview_data if r.is_valid)
    invalid_rows = total_rows - valid_rows

    return PreviewResponse(
        total_rows=total_rows,
        valid_rows=valid_rows,
        invalid_rows=invalid_rows,
        preview_data=preview_data
    )

def execute_import(request: ImportExecuteRequest) -> ImportExecuteResponse:
    success_count = 0
    failed_count = 0
    failed_rows: List[RowPreview] = []

    for row in request.rows:
        if not row.is_valid:
            failed_count += 1
            failed_rows.append(row)
            continue
            
        try:
            clean_email = row.email.strip().lower()
            
            prefix = re.sub(r'[^a-z0-9]', '', clean_email.split('@')[0].lower())
            if not prefix:
                prefix = "user"
            clean_username = prefix
            counter = 1
            while clean_username in USERS_DB:
                clean_username = f"{prefix}{counter}"
                counter += 1

            new_user = UserInDB(
                id=get_next_user_id(),
                username=clean_username,
                full_name=row.full_name,
                email=clean_email,
                phone=row.phone if row.phone else None,
                role=row.role,
                roles=[row.role],
                hashed_password=get_password_hash(row.password),
                branch=row.branch if row.branch else "Kho Tổng Hà Nội",
                is_active=True,
            )
            USERS_DB[clean_username] = new_user
            success_count += 1
        except Exception as e:
            row.errors.append(str(e))
            row.is_valid = False
            failed_rows.append(row)
            failed_count += 1

    if success_count > 0:
        save_users_db()

    return ImportExecuteResponse(
        total_processed=len(request.rows),
        success_count=success_count,
        failed_count=failed_count,
        failed_rows=failed_rows
    )

def create_template_excel() -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Users Import"
    
    headers = ["Họ và tên", "Email", "Số điện thoại", "Vai trò (Role)", "Phòng ban/Địa bàn", "Mật khẩu"]
    ws.append(headers)
    
    example_row = ["Nguyễn Văn A", "nva@congty.vn", "0912345678", "sales", "Kho Tổng Hà Nội", "123456"]
    ws.append(example_row)
    
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream

def create_error_excel(failed_rows: List[RowPreview]) -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Failed Rows"
    
    headers = ["Dòng", "Họ và tên", "Email", "Số điện thoại", "Vai trò (Role)", "Phòng ban/Địa bàn", "Mật khẩu", "Lỗi"]
    ws.append(headers)
    
    for row in failed_rows:
        ws.append([
            row.row_index,
            row.full_name,
            row.email,
            row.phone,
            row.role,
            row.branch,
            row.password,
            " | ".join(row.errors)
        ])
        
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
=======
import openpyxl
from openpyxl import Workbook
from openpyxl.utils.exceptions import InvalidFileException
from typing import List, Tuple, Dict, Any
import re
from sqlalchemy.orm import Session
from app.models.entities import UserEntity
from app.schemas.user_import import BulkImportRowResult, BulkImportPreviewResponse, BulkImportExecuteResponse
from app.core.rbac import ROLE_DETAILS

class UserBulkImportService:
    @staticmethod
    def _is_valid_email(email: str) -> bool:
        pattern = r"^[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+$"
        return re.match(pattern, email) is not None

    @staticmethod
    def _is_valid_phone(phone: str) -> bool:
        cleaned = re.sub(r'[\s\.\-\(\)]', '', phone.strip())
        if cleaned.startswith('+84'):
            cleaned = '0' + cleaned[3:]
        elif cleaned.startswith('84') and len(cleaned) == 11:
            cleaned = '0' + cleaned[2:]
        vn_phone_pattern = r'^(0[3|5|7|8|9][0-9]{8}|02[0-9]{9})$'
        return re.match(vn_phone_pattern, cleaned) is not None

    @staticmethod
    def parse_and_validate(file_content: bytes, db: Session) -> BulkImportPreviewResponse:
        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
            sheet = wb.active
        except Exception as e:
            raise ValueError(f"Không thể đọc file Excel: {str(e)}")

        # Template format: 
        # Cột A: Họ và tên
        # Cột B: Email
        # Cột C: Số điện thoại
        # Cột D: Vai trò (Role)
        # Cột E: Chi nhánh / Địa bàn
        # Cột F: Mật khẩu

        rows_result = []
        emails_in_file = set()
        phones_in_file = set()
        
        valid_count = 0
        invalid_count = 0

        # Existing emails/phones from DB
        existing_users = db.query(UserEntity.email, UserEntity.phone).all()
        existing_emails = {u.email for u in existing_users if u.email}
        existing_phones = {u.phone for u in existing_users if u.phone}
        
        valid_roles = list(ROLE_DETAILS.keys())

        # Ánh xạ vai trò tiếng Việt sang mã vai trò hệ thống
        role_map_vi = {
            "quản trị hệ thống": "admin",
            "quản trị": "admin",
            "admin": "admin",
            "quản lý kinh doanh": "sales_manager",
            "quản lý bán hàng": "sales_manager",
            "sales_manager": "sales_manager",
            "nhân viên kinh doanh": "sales",
            "kinh doanh": "sales",
            "sales": "sales",
            "sale": "sales",
            "thủ kho": "warehouse",
            "warehouse": "warehouse",
            "quản lý kho": "warehouse_manager",
            "warehouse_manager": "warehouse_manager",
            "kế toán": "accountant",
            "kế toán công nợ": "accountant",
            "accountant": "accountant",
            "nhân viên mua hàng": "purchasing",
            "mua hàng": "purchasing",
            "purchasing": "purchasing",
            "đại lý": "customer",
            "khách hàng": "customer",
            "customer": "customer",
        }

        for idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if all(v is None for v in row):
                continue
                
            full_name = str(row[0]).strip() if row[0] else ""
            email = str(row[1]).strip() if row[1] else ""
            phone = str(row[2]).strip() if row[2] else ""
            raw_role = str(row[3]).strip() if row[3] else ""
            # Chuẩn hóa vai trò từ tiếng Việt hoặc code tiếng Anh
            norm_role_key = raw_role.lower()
            role = role_map_vi.get(norm_role_key, raw_role)
            branch = str(row[4]).strip() if len(row) > 4 and row[4] else "Kho Tổng Hà Nội"
            password = str(row[5]).strip() if len(row) > 5 and row[5] else "123"

            errors = {}

            if not full_name:
                errors['full_name'] = "Bắt buộc nhập họ và tên."
            if not email:
                errors['email'] = "Bắt buộc nhập email."
            elif not UserBulkImportService._is_valid_email(email):
                errors['email'] = "Email không hợp lệ."
            else:
                if email in existing_emails:
                    errors['email'] = "Email đã tồn tại trong hệ thống."
                elif email in emails_in_file:
                    errors['email'] = "Email bị trùng lặp trong file."
                emails_in_file.add(email)

            if not phone:
                errors['phone'] = "Bắt buộc nhập số điện thoại."
            elif not UserBulkImportService._is_valid_phone(phone):
                errors['phone'] = "Số điện thoại không hợp lệ."
            else:
                cleaned_phone = re.sub(r'[\s\.\-\(\)]', '', phone)
                if cleaned_phone.startswith('+84'):
                    cleaned_phone = '0' + cleaned_phone[3:]
                elif cleaned_phone.startswith('84') and len(cleaned_phone) == 11:
                    cleaned_phone = '0' + cleaned_phone[2:]
                
                if cleaned_phone in existing_phones:
                    errors['phone'] = "Số điện thoại đã tồn tại trong hệ thống."
                elif cleaned_phone in phones_in_file:
                    errors['phone'] = "Số điện thoại bị trùng lặp trong file."
                phones_in_file.add(cleaned_phone)

            if not raw_role:
                errors['role'] = "Bắt buộc nhập vai trò."
            elif role not in valid_roles:
                errors['role'] = "Vai trò không hợp lệ. Vui lòng chọn một trong các vai trò: Nhân viên kinh doanh, Thủ kho, Quản lý kho, Kế toán, Đại lý, Quản trị hệ thống, Quản lý kinh doanh, Nhân viên mua hàng."

            if not password:
                errors['password'] = "Bắt buộc nhập mật khẩu."

            is_valid = len(errors) == 0
            if is_valid:
                valid_count += 1
            else:
                invalid_count += 1

            rows_result.append(BulkImportRowResult(
                row_index=idx,
                full_name=full_name,
                email=email,
                phone=phone,
                role=role,
                branch=branch,
                password=password,
                is_valid=is_valid,
                errors=errors
            ))
            
        return BulkImportPreviewResponse(
            rows=rows_result,
            total_rows=len(rows_result),
            valid_count=valid_count,
            invalid_count=invalid_count
        )

    @staticmethod
    def generate_template() -> bytes:
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        
        # --- Sheet 1: Dữ liệu mẫu ---
        ws = wb.active
        ws.title = "Danh sách người dùng"
        
        headers = [
            "Họ và tên (*)",
            "Email (*)",
            "Số điện thoại (*)",
            "Vai trò (*)",
            "Chi nhánh / Kho / Địa bàn",
            "Mật khẩu"
        ]
        ws.append(headers)
        
        # Mẫu dữ liệu phong phú minh họa chuẩn cho các vai trò (cho phép điền Tiếng Việt có dấu hoặc mã code tiếng Anh)
        sample_rows = [
            ["Nguyễn Văn A", "nguyenvana@gmail.com", "0912345678", "Nhân viên kinh doanh", "Kho Tổng Hà Nội", "123"],
            ["Trần Thị B", "tranthib@gmail.com", "0987654321", "Thủ kho", "Kho Chi Nhánh Đà Nẵng", "123"],
            ["Lê Văn C", "levanc@gmail.com", "0901234567", "Đại lý", "Khu vực Miền Nam", "123"],
            ["Phạm Minh D", "minhd@gmail.com", "0934567890", "Kế toán", "Trụ sở chính", "123"],
        ]
        for row in sample_rows:
            ws.append(row)

        # Định dạng Header
        header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1")
        )

        ws.row_dimensions[1].height = 28

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align
            cell.border = thin_border

        # Định dạng các dòng dữ liệu (Cột SĐT ép kiểu Text '@' để không bao giờ mất số 0)
        for row in ws.iter_rows(min_row=2, max_row=len(sample_rows) + 1, min_col=1, max_col=len(headers)):
            ws.row_dimensions[row[0].row].height = 22
            for cell in row:
                cell.font = Font(name="Calibri", size=11)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")
                if cell.column == 3:  # Số điện thoại
                    cell.number_format = "@"

        # Tự động căn chỉnh độ rộng cột chuẩn mắt (không bị che khuất)
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 6, 18)

        # --- Sheet 2: Danh mục mã vai trò & Kho hợp lệ ---
        ws_ref = wb.create_sheet(title="Hướng dẫn mã vai trò")
        ref_headers = ["Tên vai trò Tiếng Việt (Điền cột D)", "Mã vai trò tiếng Anh", "Lưu ý địa bàn / Kho bắt buộc"]
        ws_ref.append(ref_headers)

        ws_ref.row_dimensions[1].height = 26
        ref_header_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
        for col_idx in range(1, len(ref_headers) + 1):
            cell = ws_ref.cell(row=1, column=col_idx)
            cell.fill = ref_header_fill
            cell.font = header_font
            cell.alignment = header_align
            cell.border = thin_border

        ref_data = [
            ["Quản trị hệ thống", "admin", "Toàn quyền hệ thống"],
            ["Quản lý kinh doanh", "sales_manager", "Khu vực hoặc Trụ sở chính"],
            ["Nhân viên kinh doanh", "sales", "Chi nhánh hoặc địa bàn phụ trách"],
            ["Thủ kho", "warehouse", "BẮT BUỘC: Kho Tổng Hà Nội, Kho Chi Nhánh Đà Nẵng, Kho Chi Nhánh TP. Hồ Chí Minh"],
            ["Quản lý kho", "warehouse_manager", "BẮT BUỘC: Kho cụ thể"],
            ["Kế toán", "accountant", "Trụ sở chính hoặc chi nhánh"],
            ["Nhân viên mua hàng", "purchasing", "Trụ sở chính"],
            ["Đại lý", "customer", "Khu vực hoạt động của đại lý"],
        ]
        for item in ref_data:
            ws_ref.append(item)

        for row in ws_ref.iter_rows(min_row=2, max_row=len(ref_data) + 1, min_col=1, max_col=len(ref_headers)):
            ws_ref.row_dimensions[row[0].row].height = 20
            for cell in row:
                cell.font = Font(name="Calibri", size=10.5)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

        for col in ws_ref.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_ref.column_dimensions[col_letter].width = max(max_len + 6, 22)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()
>>>>>>> 96a74e689e924f38d15f302d8efbdcace34d4dd6
